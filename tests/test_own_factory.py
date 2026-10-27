"""本厂产品: style codes, archive, sales, image replacement, and library filters."""

from __future__ import annotations

import io
import json
from pathlib import Path

import fitz
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import src.catalog as catalog
from api import main as api
from src import import_jobs
from src.factory.codes import image_style_code, parse_style_code
from src.factory.colors import classify_color, split_colors
from src.factory.sizes import split_sizes
from src.schema import FashionTags
from test_review_workflow import approve, make_record, store

FIXTURES = Path(__file__).parent / "fixtures" / "own_factory"


def test_style_code_splits_base_and_customer_and_rejects_totals():
    own = parse_style_code("516083")
    assert own and own.base == "516083" and own.customer == "本厂自有版本" and own.suffix == ""
    for raw, customer in [("516083若宇", "若宇"), ("516083张学庚", "张学庚"), ("516083卡雷莱", "卡雷莱"),
                          ("516083声雨竹", "声雨竹"), ("633125*7", "*7"), ("7A933红人", "红人"), ("7A935红人", "红人")]:
        parsed = parse_style_code(raw)
        assert parsed and parsed.customer == customer and parsed.base == raw.replace(customer, "")
    upper = parse_style_code("s7529")
    assert upper and upper.base == "S7529" and upper.customer == "本厂自有版本"
    cd = parse_style_code("CD0153")
    assert cd and cd.base == "CD0153"
    assert parse_style_code("5") is None and parse_style_code("19") is None
    assert parse_style_code("款号") is None and parse_style_code("合计") is None
    assert image_style_code("516083 .JPG") == "516083"
    assert image_style_code("s7529.jpg") == "S7529"
    assert image_style_code("season/516083.jpg") == "516083"


def test_colors_and_sizes_keep_names_and_split_chinese_separators():
    colors = split_colors("驼灰，玛瑙红、印象米")
    assert [item["name"] for item in colors] == ["驼灰", "玛瑙红", "印象米"]
    assert [item["family"] for item in colors] == ["灰", "红", "米白"]
    assert classify_color("深花啡") == "棕咖"
    assert classify_color("亮紫薯") == "紫"
    assert classify_color("落日桔") == "橙"
    assert classify_color("亮苔绿") == "绿"
    assert classify_color("上青+白") == "蓝"
    assert [item["name"] for item in split_colors("上青+白")] == ["上青+白"]
    assert classify_color("没见过的颜色") == "其他"
    sizes = split_sizes("95/S、100/M,115/2XL")
    assert [(item["label"], item["number"], item["letter"]) for item in sizes] == [
        ("95/S", "95", "S"), ("100/M", "100", "M"), ("115/2XL", "115", "2XL"),
    ]


def _client(monkeypatch):
    monkeypatch.setattr(import_jobs, "submit_job", import_jobs.run_job)
    return TestClient(api.app)


def _upload(client, name, data, category, content_type="image/jpeg"):
    response = client.post(
        "/api/import/jobs?backend=mock",
        files=[("files", (name, data, content_type)), ("source_categories", (None, category))],
    )
    assert response.status_code == 202, response.text
    return client.get(f"/api/import/jobs/{response.json()['id']}").json()


def _jpeg(color, size=(240, 320)):
    buf = io.BytesIO()
    Image.new("RGB", size, color).save(buf, format="JPEG", quality=90)
    return buf.getvalue()


def _sheet(path: Path, headers, rows):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    book.save(path)


def test_archive_and_sales_samples_upsert_and_skip_totals(store):
    client = TestClient(api.app)
    archive = client.post(
        "/api/import/style-archive",
        files={"file": ("style_archive_516083.xlsx", (FIXTURES / "style_archive_516083.xlsx").read_bytes())},
    )
    assert archive.status_code == 200, archive.text
    body = archive.json()
    assert body["variant_count"] == 5
    assert body["skipped_rows"] >= 1
    assert any("系列" in warning for warning in body["warnings"])
    assert body["archive_without_image"] == ["516083"]
    customers = {row[0] for row in _variants(store)}
    assert customers == {"本厂自有版本", "若宇", "张学庚", "卡雷莱", "声雨竹"}
    assert all(not code.startswith("KG") for code in _variant_codes())

    full = client.post(
        "/api/import/style-archive",
        files={"file": ("style_archive_full_sample.xlsx", (FIXTURES / "style_archive_full_sample.xlsx").read_bytes())},
    )
    assert full.status_code == 200, full.text
    assert full.json()["variant_count"] == 19
    codes = _variant_codes()
    assert "633125*7" in codes and "7A933红人" in codes and "S7529" in codes
    assert "KG516083" not in codes
    customers = {row[0] for row in _variants(store)}
    assert "*7" in customers and "红人" in customers
    assert not any("红人服饰" in name for name in customers)

    sales = client.post(
        "/api/import/sales",
        files={"file": ("sales_516083.xlsx", (FIXTURES / "sales_516083.xlsx").read_bytes())},
    )
    assert sales.status_code == 200, sales.text
    report = sales.json()
    assert report["parsed_quantity"] == 3411
    assert report["parsed_amount"] == 509145
    assert report["stored_quantity"] == 3411
    assert report["stored_amount"] == 509145
    assert report["skipped_rows"] >= 2
    again = client.post(
        "/api/import/sales",
        files={"file": ("sales_516083.xlsx", (FIXTURES / "sales_516083.xlsx").read_bytes())},
    )
    assert again.json()["updated"] >= 1
    assert again.json()["stored_quantity"] == 3411
    assert again.json()["stored_amount"] == 509145
    again_archive = client.post(
        "/api/import/style-archive",
        files={"file": ("style_archive_516083.xlsx", (FIXTURES / "style_archive_516083.xlsx").read_bytes())},
    )
    assert again_archive.json()["variant_count"] == 5
    assert sum(1 for code in _variant_codes() if code.startswith("516083")) == 5


def test_edge_archive_without_image_and_chinese_separators(store):
    path = store / "edge.xlsx"
    _sheet(path, ["季节", "大货款号", "颜色", "尺码", "原料名称", "针型", "一级品类"], [
        ["2025春", "NOIMG88", "驼灰，玛瑙红、印象米", "95/S、100/M", "羊毛", "12针", "背心"],
        ["", "5", "", "", "", "", ""],
    ])
    client = TestClient(api.app)
    report = client.post("/api/import/style-archive", files={"file": ("edge.xlsx", path.read_bytes())}).json()
    assert report["variant_count"] == 1
    assert "NOIMG88" in report["archive_without_image"]
    colors = split_colors("驼灰，玛瑙红、印象米")
    stored = json.loads(_one("SELECT colors_json FROM style_variants WHERE variant_code='NOIMG88'"))
    assert stored == colors
    sizes = json.loads(_one("SELECT sizes_json FROM style_variants WHERE variant_code='NOIMG88'"))
    assert [item["letter"] for item in sizes] == ["S", "M"]


def test_image_upload_folder_case_replacement_and_filters(store, monkeypatch):
    client = _client(monkeypatch)
    first = _jpeg("navy")
    job = _upload(client, "season/516083 .JPG", first, "本厂产品")
    assert job["succeeded"] == 1
    record = catalog.load_catalog()[0]
    assert record.source_category == "本厂产品"
    assert record.style_code == "516083"
    assert record.notes.款号 == "516083"
    assert record.source.location == "season/516083 .JPG"
    assert record.pending_image_path is None

    upper = _upload(client, "s7529.JPG", _jpeg("green"), "本厂产品")
    assert upper["succeeded"] == 1
    assert any(item.style_code == "S7529" for item in catalog.load_catalog())

    same = _upload(client, "516083.jpg", first, "本厂产品")
    assert same["succeeded"] == 1
    assert "替换图" in same["message"]
    assert len(catalog.load_catalog()) == 2
    repeated = next(item for item in catalog.load_catalog() if item.style_code == "516083")
    assert repeated.pending_image_path
    assert Path(repeated.image_path).read_bytes() == Path(repeated.pending_image_path).read_bytes()

    replacement = _upload(client, "516083.jpg", _jpeg("red"), "本厂产品")
    assert replacement["succeeded"] == 1
    assert "替换图" in replacement["message"]
    pending = next(item for item in catalog.load_catalog() if item.style_code == "516083")
    assert pending.review_status == "pending"
    assert pending.pending_image_path
    assert Path(pending.image_path).read_bytes() != Path(pending.pending_image_path).read_bytes()
    assert len(catalog.load_catalog()) == 2

    archive = client.post(
        "/api/import/style-archive",
        files={"file": ("a.xlsx", (FIXTURES / "style_archive_516083.xlsx").read_bytes())},
    )
    assert archive.status_code == 200
    sales = client.post(
        "/api/import/sales",
        files={"file": ("s.xlsx", (FIXTURES / "sales_516083.xlsx").read_bytes())},
    )
    assert sales.json()["stored_quantity"] == 3411
    detail = client.get(f"/api/records/{pending.id}").json()
    assert len(detail["variants"]) == 5
    assert detail["sales"]["total_quantity"] == 3411
    assert detail["sales"]["total_amount"] == 509145
    assert {row["customer"] for row in detail["variants"]} == {"本厂自有版本", "若宇", "张学庚", "卡雷莱", "声雨竹"}
    assert detail["library_category"] == "上衣>套衫"
    assert detail["sleeve_length"] == "长袖"
    assert detail["fabric"] == "针织"
    assert detail["archive_category"] == "长袖套衫"
    assert {row["category"] for row in detail["variants"]} == {"长袖套衫"}
    assert all("KG" not in row["variant_code"] for row in detail["variants"])
    assert detail["has_pending_image"] and detail["pending_image_url"]
    assert client.get(detail["pending_image_url"]).status_code == 200

    library = client.get("/api/catalog", params={"source_category": "本厂产品", "sort": "sales"}).json()
    assert library["items"][0]["style_code"] == "516083"
    assert library["items"][0]["sales_quantity"] == 3411
    for params, expected in [
        ({"source_category": "本厂产品", "archive_season": "2025秋"}, True),
        ({"source_category": "本厂产品", "color_family": "紫"}, True),
        ({"source_category": "本厂产品", "size": "M"}, True),
        ({"source_category": "本厂产品", "customer": "若宇"}, True),
        ({"source_category": "本厂产品", "category": "上衣", "subcategory": "套衫", "sleeve_length": "长袖"}, True),
        ({"group": "上衣", "source_category": "本厂产品"}, True),
        ({"group": "下装", "source_category": "本厂产品"}, False),
        ({"source_category": "网图趋势"}, False),
    ]:
        found = client.get("/api/catalog", params=params).json()["items"]
        assert (pending.id in {item["id"] for item in found}) is expected, params

    hits = client.get("/api/search/text", params={"q": "卡雷莱", "include_pending": "true", "source_category": "本厂产品"}).json()
    assert pending.id in {item["id"] for item in hits["items"]}
    other = client.get("/api/search/text", params={"q": "516083", "include_pending": "true", "source_category": "市调拍照"}).json()
    assert pending.id not in {item["id"] for item in other["items"]}

    approved = client.post(f"/api/records/{pending.id}/review", json={
        "action": "approve", "reviewer": "审核员", "expected_revision": detail["revision"],
        "tags": FashionTags(品类="毛衣", 季节="秋").model_dump(),
    })
    assert approved.status_code == 200, approved.text
    swapped = catalog.load_catalog()
    current = next(item for item in swapped if item.id == pending.id)
    assert current.review_status == "approved"
    assert current.pending_image_path is None
    assert current.image_history and Path(current.image_history[0]["image_path"]).is_file()
    history = client.get(f"/api/records/{pending.id}").json()["image_history"]
    assert history and client.get(history[0]["url"]).status_code == 200

    second = _upload(client, "516083.jpg", _jpeg("white"), "本厂产品")
    assert "替换图" in second["message"]
    current = next(item for item in catalog.load_catalog() if item.id == pending.id)
    assert current.review_status == "pending" and current.pending_image_path
    kept = client.post(f"/api/records/{pending.id}/review", json={
        "action": "keep_image", "reviewer": "审核员", "note": "保留原图", "expected_revision": current.revision,
    })
    assert kept.status_code == 200, kept.text
    restored = next(item for item in catalog.load_catalog() if item.id == pending.id)
    assert restored.review_status == "approved"
    assert restored.pending_image_path is None
    assert restored.image_path == current.image_path


def test_duplicate_style_code_cannot_be_reassigned(store, monkeypatch):
    client = _client(monkeypatch)
    _upload(client, "516083.jpg", _jpeg("navy"), "本厂产品")
    _upload(client, "516083.jpg", _jpeg("olive"), "网图趋势")
    records = catalog.load_catalog()
    other = next(item for item in records if item.source_category == "网图趋势")
    assert other.style_code == ""
    response = client.post("/api/records/source-category", json={"ids": [other.id], "source_category": "本厂产品"})
    assert response.status_code == 409
    assert "款号" in response.json()["detail"]
    assert next(item for item in catalog.load_catalog() if item.id == other.id).source_category == "网图趋势"


def test_office_import_keeps_non_factory_categories(store, monkeypatch):
    client = _client(monkeypatch)
    image = io.BytesIO()
    Image.new("RGB", (240, 320), "navy").save(image, format="PNG")
    image.seek(0)
    ppt = Presentation()
    ppt.slides.add_slide(ppt.slide_layouts[6]).shapes.add_picture(image, Inches(0.2), Inches(0.2), Inches(2), Inches(3))
    ppt_data = io.BytesIO()
    ppt.save(ppt_data)

    png = store / "sheet.png"
    Image.new("RGB", (240, 320), "teal").save(png)
    book = Workbook()
    sheet = book.active
    sheet.add_image(XLImage(str(png)), "A1")
    xlsx = io.BytesIO()
    book.save(xlsx)

    doc = fitz.open()
    page = doc.new_page(width=400, height=560)
    page.insert_image(fitz.Rect(40, 40, 280, 360), stream=image.getvalue())
    pdf_data = doc.tobytes()
    doc.close()

    response = client.post("/api/import/jobs?backend=mock", files=[
        ("files", ("lookbook.pptx", ppt_data.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")),
        ("files", ("sheet.xlsx", xlsx.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ("files", ("page.pdf", pdf_data, "application/pdf")),
        ("source_categories", (None, "网图趋势")),
        ("source_categories", (None, "品牌资料包")),
        ("source_categories", (None, "市调拍照")),
    ])
    assert response.status_code == 202, response.text
    body = client.get(f"/api/import/jobs/{response.json()['id']}").json()
    assert body["failed"] == 0, body
    assert body["succeeded"] == 3
    records = catalog.load_catalog()
    assert {item.source_category for item in records} == {"网图趋势", "品牌资料包", "市调拍照"}
    assert all(item.style_code == "" for item in records)
    assert client.get("/api/catalog", params={"source_category": "网图趋势"}).json()["count"] == 1
    plain = client.post("/api/import/jobs?backend=mock", files={"files": ("again.pptx", ppt_data.getvalue())})
    assert plain.status_code == 202, plain.text
    finished = client.get(f"/api/import/jobs/{plain.json()['id']}").json()
    assert finished["succeeded"] == 1, finished
    assert any(item.source_category == "未分类" for item in catalog.load_catalog())


def test_legacy_rows_gain_unclassified_without_revision_bump(store):
    record = make_record(store)
    catalog.save_catalog([record])
    approved = approve(record)
    with catalog._database() as conn:
        payload = json.loads(conn.execute("SELECT payload FROM records WHERE id=?", (record.id,)).fetchone()[0])
        payload.pop("source_category", None)
        payload.pop("style_code", None)
        conn.execute("UPDATE records SET payload=? WHERE id=?", (json.dumps(payload, ensure_ascii=False), record.id))
        conn.execute("DELETE FROM migrations WHERE name='source_category_v1'")
    loaded = catalog.load_catalog()[0]
    assert loaded.source_category == "未分类"
    assert loaded.style_code == ""
    assert loaded.revision == approved.revision
    assert loaded.review_status == "approved"
    again = catalog.review_record(record.id, action="resubmit", reviewer="甲", expected_revision=approved.revision)
    assert again.review_status == "pending"


def test_placeholder_key_is_flagged_without_removing_a_real_key(store, monkeypatch):
    from src.tag.tagger import has_openai_key, openai_key_is_placeholder

    monkeypatch.setenv("OPENAI_API_KEY", "sk-...")
    assert openai_key_is_placeholder()
    assert has_openai_key()
    monkeypatch.setenv("OPENAI_API_KEY", "sk-real-key-value-0123456789")
    assert not openai_key_is_placeholder()
    client = TestClient(api.app)
    monkeypatch.setenv("OPENAI_API_KEY", "sk-...")
    health = client.get("/api/health").json()
    assert health["openai_placeholder"] is True
    assert health["openai_ready"] is False


def test_wrong_table_and_image_slot_only_change_the_wording(store, monkeypatch):
    client = _client(monkeypatch)
    sales = (FIXTURES / "sales_516083.xlsx").read_bytes()
    archive = (FIXTURES / "style_archive_516083.xlsx").read_bytes()
    swapped = client.post("/api/import/style-archive", files={"file": ("sales.xlsx", sales)})
    assert swapped.status_code == 422
    assert "可能传错了表" in swapped.json()["detail"] and "销量表" in swapped.json()["detail"]
    swapped = client.post("/api/import/sales", files={"file": ("archive.xlsx", archive)})
    assert swapped.status_code == 422
    assert "可能传错了表" in swapped.json()["detail"] and "款式档案" in swapped.json()["detail"]
    job = client.post("/api/import/jobs?backend=mock", files=[
        ("files", ("sales_516083.xlsx", sales, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")),
        ("source_categories", (None, "本厂产品")),
    ])
    assert job.status_code == 202, job.text
    finished = client.get(f"/api/import/jobs/{job.json()['id']}").json()
    assert "这像是销量表，请用下方对应入口" in finished["message"]
    assert finished["succeeded"] == 0
    assert client.get("/api/catalog").json()["count"] == 0


def test_library_search_finds_customer_and_color_names(store, monkeypatch):
    client = _client(monkeypatch)
    _upload(client, "516083.jpg", _jpeg("navy"), "本厂产品")
    _upload(client, "later.jpg", _jpeg("red"), "网图趋势")
    archive = client.post("/api/import/style-archive", files={"file": ("a.xlsx", (FIXTURES / "style_archive_516083.xlsx").read_bytes())})
    assert archive.status_code == 200, archive.text
    sales = client.post("/api/import/sales", files={"file": ("s.xlsx", (FIXTURES / "sales_516083.xlsx").read_bytes())})
    assert sales.status_code == 200, sales.text

    by_customer = client.get("/api/catalog", params={"q": "若宇", "source_category": "本厂产品"}).json()
    assert by_customer["count"] == 1
    assert by_customer["items"][0]["style_code"] == "516083"
    assert by_customer["items"][0]["variant_count"] == 5
    assert "灰" in by_customer["items"][0]["color_families"]
    by_color = client.get("/api/catalog", params={"q": "驼灰", "source_category": "本厂产品"}).json()
    assert [item["style_code"] for item in by_color["items"]] == ["516083"]

    listed = client.get("/api/catalog", params={"sort": "imported"}).json()
    assert listed["items"][0]["source_category"] == "网图趋势"
    plain = client.get("/api/catalog").json()
    assert [item["source_category"] for item in plain["items"]] == ["本厂产品", "网图趋势"]
    amounts = client.get("/api/catalog", params={"sort": "amount", "source_category": "本厂产品"}).json()
    assert amounts["items"][0]["sales_amount"] == 509145
    assert amounts["color_counts"]["灰"] == 1
    assert amounts["color_counts"]["紫"] == 1


def test_review_reject_still_removes_factory_image_from_library(store, monkeypatch):
    client = _client(monkeypatch)
    _upload(client, "516083.jpg", _jpeg("navy"), "本厂产品")
    item = client.get("/api/reviews").json()["items"][0]
    rejected = client.post(f"/api/records/{item['id']}/review", json={
        "action": "reject", "reviewer": "审核员", "note": "不是这款", "expected_revision": item["revision"],
    })
    assert rejected.status_code == 200, rejected.text
    assert client.get("/api/catalog").json()["count"] == 0
    assert client.get("/api/reviews", params={"status": "rejected"}).json()["count"] == 1


def _variants(store):
    del store
    with catalog._database() as conn:
        return list(conn.execute("SELECT customer, variant_code FROM style_variants"))


def _variant_codes():
    with catalog._database() as conn:
        return {row[0] for row in conn.execute("SELECT variant_code FROM style_variants")}


def _one(sql):
    with catalog._database() as conn:
        return conn.execute(sql).fetchone()[0]
