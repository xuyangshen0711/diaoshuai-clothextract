from __future__ import annotations

import io
import json
from concurrent.futures import ThreadPoolExecutor

import pytest
from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

import src.catalog as catalog
import src.extract.office as office
import src.index.embeddings as embeddings
from api import main as api
from src import import_jobs
from src.schema import FashionTags, ImageRecord, SourceMeta
from src.search.query import search_by_image, search_by_text
from src.tag.tagger import tag_records


@pytest.fixture
def store(tmp_path, monkeypatch):
    for name, path in {
        "DATA_DIR": tmp_path, "RAW_DIR": tmp_path / "raw", "EXTRACTED_DIR": tmp_path / "extracted",
        "PROCESSED_DIR": tmp_path / "processed", "CATALOG_PATH": tmp_path / "processed/catalog.json",
        "DATABASE_PATH": tmp_path / "processed/catalog.sqlite3", "INDEX_DIR": tmp_path / "processed/index",
    }.items():
        monkeypatch.setattr(catalog, name, path)
    monkeypatch.setattr(embeddings, "INDEX_DIR", catalog.INDEX_DIR)
    monkeypatch.setattr(api, "RAW_DIR", catalog.RAW_DIR)
    monkeypatch.setattr(office, "EXTRACTED_DIR", catalog.EXTRACTED_DIR)
    monkeypatch.setenv("EMBEDDING_BACKEND", "local")
    catalog.ensure_data_dirs()
    return tmp_path


def make_record(store, name="sweater", color="navy"):
    path = store / "extracted" / f"{name}.png"
    Image.new("RGB", (240, 320), color).save(path)
    record = ImageRecord(image_path=str(path), source=SourceMeta(
        source_file="lookbook.pptx", source_type="pptx", location="slide 1"))
    return tag_records([record], backend="mock")[0]


def approve(record, tags=None, reviewer="设计师甲"):
    return catalog.review_record(record.id, action="approve", reviewer=reviewer,
                                 expected_revision=record.revision, tags=tags)


def test_legacy_migration_is_pending_idempotent_and_keeps_source(store):
    record = make_record(store)
    record.tags_source = "manual"
    record.manually_edited = True
    raw = json.dumps([record.model_dump()], ensure_ascii=False)
    catalog.CATALOG_PATH.write_text(raw)
    first = catalog.load_catalog()
    assert len(first) == 1 and first[0].review_status == "pending"
    assert first[0].proposed_tags == record.tags
    assert len(catalog.load_catalog()) == 1
    assert catalog.CATALOG_PATH.read_text() == raw
    assert catalog.DATABASE_PATH.exists()
    assert [e["action"] for e in catalog.review_history(record.id)] == ["legacy_imported"]


def test_modify_then_approve_preserves_ai_and_audit(store):
    record = make_record(store)
    catalog.save_catalog([record])
    final_tags = record.tags.model_copy(update={"品类": "外套"})
    result = approve(record, final_tags)
    assert result.review_status == "approved" and result.tags.品类 == "外套"
    assert result.ai_tags.品类 == "毛衣" and result.proposed_tags is None
    assert result.manually_edited and result.reviewer == "设计师甲" and result.reviewed_at
    persisted = catalog.load_catalog()[0]
    assert persisted == result
    history = catalog.review_history(record.id)
    assert history[0]["before"]["ai_tags"]["品类"] == "毛衣"
    assert history[0]["after"]["tags"]["品类"] == "外套"
    assert history[0]["actor"] == "设计师甲"


def test_retag_keeps_human_result_and_requires_new_approval(store, monkeypatch):
    record = make_record(store)
    catalog.save_catalog([record])
    approved = approve(record, record.tags.model_copy(update={"品类": "外套"}))
    assert tag_records([approved], backend="mock")[0].review_status == "approved"
    tag_records([approved], backend="mock", overwrite=True)
    assert approved.review_status == "pending"
    assert approved.tags.品类 == "外套" and approved.proposed_tags.品类 == "毛衣"
    catalog.save_catalog([approved])
    assert not catalog.approved_records(catalog.load_catalog())
    final = approve(approved)
    assert final.tags.品类 == "毛衣" and final.tags_source == "mock"
    assert len([e for e in catalog.review_history(record.id) if e["action"] == "approve"]) == 2


def test_reject_reason_and_resubmit(store):
    record = make_record(store)
    catalog.save_catalog([record])
    with pytest.raises(ValueError, match="原因"):
        catalog.review_record(record.id, action="reject", reviewer="甲", expected_revision=record.revision)
    rejected = catalog.review_record(record.id, action="reject", reviewer="甲",
        expected_revision=record.revision, note="不是服装图")
    assert rejected.review_status == "rejected" and rejected.review_note == "不是服装图"
    pending = catalog.review_record(record.id, action="resubmit", reviewer="乙",
        expected_revision=rejected.revision)
    assert pending.review_status == "pending"
    assert approve(pending).review_status == "approved"


def test_concurrent_review_only_one_decision_wins(store):
    record = make_record(store)
    catalog.save_catalog([record])
    def decide():
        try:
            approve(record)
            return "approved"
        except catalog.CatalogConflict:
            return "conflict"
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(lambda _: decide(), range(2))) == ["approved", "conflict"]
    assert len([e for e in catalog.review_history(record.id) if e["action"] == "approve"]) == 1


def test_stale_catalog_write_rolls_back_all_changes(store):
    first, second = make_record(store), make_record(store, "dress")
    catalog.save_catalog([first, second])
    stale = catalog.load_catalog()
    approve(first)
    stale[1].tags.品类 = "裤装"
    with pytest.raises(catalog.CatalogConflict):
        catalog.save_catalog([stale[1], stale[0]])
    records = catalog.load_catalog()
    assert records[0].review_status == "approved" and records[1].tags.品类 == "连衣裙"


def test_catalog_includes_pending_while_searches_require_approval_and_refresh_index(store):
    pending, rejected, good = (make_record(store, name, color) for name, color in
        [("sweater", "navy"), ("dress", "red"), ("pants", "black")])
    catalog.save_catalog([pending, rejected, good])
    catalog.review_record(rejected.id, action="reject", reviewer="甲", note="无效图", expected_revision=rejected.revision)
    good = approve(good)
    client = TestClient(api.app)
    assert [x["id"] for x in client.get("/api/catalog").json()["items"]] == [pending.id, good.id]
    records = catalog.load_catalog()
    assert all(r.review_status == "approved" for r, _ in search_by_text(records, "裤装"))
    image_hits = search_by_image(records, good.image_path)
    assert [r.id for r, _ in image_hits] == [good.id]
    _, ids, _ = embeddings.load_index()
    assert ids == [good.id]
    catalog.review_record(good.id, action="resubmit", reviewer="甲", expected_revision=good.revision)
    assert client.get("/api/catalog").json()["count"] == 2
    assert search_by_image(catalog.load_catalog(), good.image_path) == []
    assert search_by_text(catalog.load_catalog(), "裤装") == []
    # Old index entries cannot leak into search even before an explicit rebuild.
    pending = approve(catalog.load_catalog()[0])
    assert [r.id for r, _ in search_by_image(catalog.load_catalog(), pending.image_path)] == [pending.id]


def test_api_import_approve_history_and_conflicts(store, monkeypatch):
    image = io.BytesIO()
    Image.new("RGB", (240, 320), "navy").save(image, format="PNG")
    image.seek(0)
    ppt = Presentation()
    ppt.slides.add_slide(ppt.slide_layouts[6]).shapes.add_picture(image, Inches(1), Inches(1))
    data = io.BytesIO()
    ppt.save(data)
    monkeypatch.setattr(import_jobs, "submit_job", import_jobs.run_job)
    client = TestClient(api.app)
    response = client.post("/api/import/jobs?backend=mock", files={"files": ("lookbook.pptx", data.getvalue())})
    assert response.status_code == 202, response.text
    assert client.get(f'/api/import/jobs/{response.json()["id"]}').json()["succeeded"] == 1
    assert client.get("/api/catalog").json()["count"] == 1
    queue = client.get("/api/reviews").json()
    item = queue["items"][0]
    assert queue["counts"]["pending"] == 1
    body = dict(action="approve", reviewer="设计师", expected_revision=item["revision"],
                tags=FashionTags(品类="毛衣", 季节="秋").model_dump())
    assert client.post(f'/api/records/{item["id"]}/review', json=body).status_code == 200
    assert client.post(f'/api/records/{item["id"]}/review', json=body).status_code == 409
    assert client.get("/api/catalog").json()["count"] == 1
    assert client.get(f'/api/records/{item["id"]}/history').json()["items"][0]["actor"] == "设计师"
    assert client.get(item["image_url"]).status_code == 200
    # Database and original office files must not be exposed as static media.
    assert client.get("/media/processed/catalog.sqlite3").status_code == 404


@pytest.mark.parametrize("change", [dict(reviewer="  "), dict(tags={"品类": "不存在", "季节": "秋"}),
    dict(action="reject", note=" ")])
def test_api_invalid_decisions_do_not_change_state(store, change):
    record = make_record(store)
    catalog.save_catalog([record])
    client = TestClient(api.app)
    body = dict(action="approve", reviewer="甲", expected_revision=record.revision)
    body.update(change)
    assert client.post(f"/api/records/{record.id}/review", json=body).status_code == 422
    assert catalog.load_catalog()[0].review_status == "pending"
    assert len(catalog.review_history(record.id)) == 1


def test_api_retag_cannot_bypass_review(store, monkeypatch):
    record = make_record(store)
    catalog.save_catalog([record])
    record = approve(record, record.tags.model_copy(update={"品类": "外套"}))
    monkeypatch.setattr(import_jobs, "submit_job", import_jobs.run_job)
    client = TestClient(api.app)
    response = client.post("/api/tag/jobs", json={"backend": "mock", "overwrite": True})
    assert response.status_code == 202
    record = catalog.load_catalog()[0]
    assert record.review_status == "pending" and record.tags.品类 == "外套"
    assert client.get("/api/catalog").json()["count"] == 1


def test_library_uses_pending_proposal_for_display_and_every_filter_without_changing_history(store):
    record = make_record(store)
    record.source.source_file = "冬季款式.xlsx"
    record.source.source_type = "xlsx"
    record.source.uploaded_at = "2026-10-20T03:00:00+00:00"
    catalog.save_catalog([record])
    record = approve(record, record.tags.model_copy(update={"品类": "外套"}))
    tag_records([record], backend="mock", overwrite=True)
    record.proposed_tags = record.proposed_tags.model_copy(update={
        "品类": "裤装", "颜色主色调": "红色", "风格": "复古", "一句话描述": "本轮新的红色裤装",
    })
    catalog.save_catalog([record])
    before = catalog.get_by_id(catalog.load_catalog(), record.id).model_dump()
    history_count = len(catalog.review_history(record.id))
    client = TestClient(api.app)
    visible = client.get("/api/catalog", params={"review_status": "pending", "group": "下装",
        "category": "裤装", "color": "红色", "style": "复古", "q": "本轮新的"}).json()
    assert visible["count"] == 1
    item = visible["items"][0]
    assert item["display_tags"]["品类"] == "裤装" and item["tags"]["品类"] == "外套"
    assert item["source"]["source_file"] == "冬季款式.xlsx"
    assert item["source"]["source_type"] == "xlsx"
    assert item["source"]["uploaded_at"] == "2026-10-20T03:00:00+00:00"
    assert client.get("/api/catalog?category=外套").json()["count"] == 0
    assert client.get("/api/catalog?review_status=approved").json()["count"] == 0
    assert client.get("/api/catalog?review_status=rejected").status_code == 422
    assert catalog.get_by_id(catalog.load_catalog(), record.id).model_dump() == before
    assert len(catalog.review_history(record.id)) == history_count
    approved = approve(record, record.tags.model_copy(update={"品类": "衬衫"}))
    result = client.get("/api/catalog?review_status=approved&category=衬衫").json()["items"][0]
    assert result["display_tags"]["品类"] == "衬衫"
    assert result["source"] == item["source"]
    pending = catalog.review_record(approved.id, action="resubmit", reviewer="甲", expected_revision=approved.revision)
    catalog.review_record(pending.id, action="reject", reviewer="甲", note="图片重复",
                          expected_revision=pending.revision)
    assert client.get("/api/catalog").json()["count"] == 0


def test_library_hides_untagged_images_and_exposes_legacy_import_time(store):
    record = make_record(store)
    untagged = ImageRecord(image_path=record.image_path + ".untagged", source=record.source)
    catalog.save_catalog([record, untagged])
    item = TestClient(api.app).get("/api/catalog").json()["items"]
    assert len(item) == 1 and item[0]["id"] == record.id
    assert item[0]["source"]["uploaded_at"] is None
    assert item[0]["source"]["extracted_at"] == record.source.extracted_at


def test_filtered_image_search_finds_candidates_beyond_global_top_k(store):
    import numpy as np

    class TestBackend:
        name, dim = "test-ranking", 2
        def encode_image(self, path):
            return np.array([1, 0], dtype=np.float32)
        def encode_record(self, record):
            return np.array([0, 1] if record.id == "target" else [1, 0], dtype=np.float32)

    base = make_record(store)
    records = [base.model_copy(update={"id": f"item-{i}", "review_status": "approved"}) for i in range(100)]
    target = base.model_copy(update={"id": "target", "review_status": "approved"})
    records.append(target)
    hits = search_by_image(records, base.image_path, top_k=1, filtered=[target], backend=TestBackend())
    assert [r.id for r, _ in hits] == ["target"]


def test_invalid_legacy_migration_rolls_back_and_can_retry(store):
    record = make_record(store)
    catalog.CATALOG_PATH.write_text(json.dumps([record.model_dump(), {"id": "broken"}]))
    with pytest.raises(ValueError):
        catalog.load_catalog()
    catalog.CATALOG_PATH.write_text(json.dumps([record.model_dump()]))
    assert len(catalog.load_catalog()) == 1
    assert len(catalog.review_history(record.id)) == 1
