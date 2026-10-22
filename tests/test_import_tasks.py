from __future__ import annotations

import io
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor
from threading import Event

from fastapi.testclient import TestClient
from PIL import Image
from pptx import Presentation
from pptx.util import Inches

from test_review_workflow import store, make_record, approve
from src import catalog, import_jobs as tasks
from src.schema import FashionTags
from src.tag.tagger import MockTagger
from api import main as api


def fixture_sources(store, monkeypatch, count=3):
    paths = []
    for i in range(count):
        path = store / 'raw' / f'file-{i}.pdf'
        path.write_bytes(b'fixture')
        paths.append(path)
    def extract(path, out_dir):
        return [make_record(store, 'sweater-' + path.stem)]
    monkeypatch.setattr(tasks, 'extract_from_file', extract)
    return paths


def test_partial_task_retries_only_failed_items(store, monkeypatch):
    paths = fixture_sources(store, monkeypatch)
    calls = []
    class OnceFailTagger(MockTagger):
        def tag(self, image_path, context=''):
            calls.append(image_path)
            if 'file-1' in image_path and calls.count(image_path) == 1:
                raise TimeoutError('Sensitive upstream error should not be exposed')
            return super().tag(image_path, context)
    monkeypatch.setattr(tasks, 'get_tagger', lambda _: OnceFailTagger())
    job = tasks.create_import(paths, 'mock')
    tasks.run_job(job['id'])
    public = tasks.public_job(tasks.get_job(job['id']))
    assert public['status'] == 'partial' and public['succeeded'] == 2 and public['failed'] == 1
    assert 'Sensitive' not in str(public) and len(catalog.load_catalog()) == 2
    tasks.retry_job(job['id'])
    tasks.run_job(job['id'])
    result = tasks.public_job(tasks.get_job(job['id']))
    assert result['status'] == 'completed' and result['succeeded'] == 3
    assert len(calls) == 4 and len(catalog.load_catalog()) == 3
    assert all(r.source.batch_id == job['id'] for r in catalog.load_catalog())
    assert all(r.review_status == 'pending' for r in catalog.load_catalog())
    with TestClient(api.app) as client:
        assert client.get('/api/catalog').json()['count'] == 3
        assert client.post(f"/api/import/jobs/{job['id']}/retry").status_code == 409


def test_live_progress_does_not_block_catalog(store, monkeypatch):
    paths = fixture_sources(store, monkeypatch, count=1)
    entered, release = Event(), Event()
    class BlockingTagger(MockTagger):
        def tag(self, path, context=''):
            entered.set()
            assert release.wait(5)
            return super().tag(path, context)
    monkeypatch.setattr(tasks, 'get_tagger', lambda _: BlockingTagger())
    job = tasks.create_import(paths, 'mock')
    with ThreadPoolExecutor(max_workers=1) as pool:
        future = pool.submit(tasks.run_job, job['id'])
        try:
            assert entered.wait(5)
            public = tasks.public_job(tasks.get_job(job['id']))
            assert public['status'] == 'tagging' and public['total'] == 1 and public['processed'] == 0
            assert TestClient(api.app).get('/api/catalog').status_code == 200
        finally:
            release.set()
        future.result(timeout=5)
    assert len(catalog.load_catalog()) == 1


def test_restart_marks_interrupted_and_requires_explicit_retry(store, monkeypatch):
    paths = fixture_sources(store, monkeypatch, 1)
    job = tasks.create_import(paths, 'mock')
    with catalog._database() as conn:
        job['status'] = 'extracting'
        tasks._put(conn, job)
    tasks.recover_interrupted()
    recovered = tasks.get_job(job['id'])
    assert recovered['status'] == 'interrupted' and catalog.load_catalog() == []
    tasks.retry_job(job['id'])
    tasks.run_job(job['id'])
    assert tasks.get_job(job['id'])['status'] == 'completed'


def test_job_endpoint_upload_source_detail_and_download(store, monkeypatch):
    submitted = []
    monkeypatch.setattr(tasks, 'submit_job', submitted.append)
    ppt = Presentation()
    image = io.BytesIO()
    Image.new('RGB', (240, 320), 'navy').save(image, format='PNG')
    image.seek(0)
    ppt.slides.add_slide(ppt.slide_layouts[6]).shapes.add_picture(image, Inches(1), Inches(1))
    payload = io.BytesIO(); ppt.save(payload)
    client = TestClient(api.app)
    response = client.post('/api/import/jobs?backend=mock', files={'files': ('new.pptx', payload.getvalue())})
    assert response.status_code == 202 and response.json()['status'] == 'queued'
    job_id = response.json()['id']
    assert submitted == [job_id] and catalog.load_catalog() == []
    tasks.run_job(job_id)
    assert client.get(f'/api/import/jobs/{job_id}').json()['succeeded'] == 1
    record = catalog.load_catalog()[0]
    detail = client.get(f'/api/records/{record.id}').json()
    assert detail['source']['batch_id'] == job_id and detail['source_url']
    assert detail['source']['uploaded_at'] == response.json()['created_at']
    assert client.get(detail['source_url']).content == payload.getvalue()
    downloaded = client.get(detail['download_url'])
    assert downloaded.status_code == 200 and 'attachment' in downloaded.headers['content-disposition']
    queue = client.get('/api/reviews', params={'batch': job_id}).json()
    assert queue['count'] == 1 and queue['facets']['batches'][0]['id'] == job_id
    assert client.get('/api/import/jobs/missing').status_code == 404


def test_quota_error_stops_batch_and_redacts_credentials(store, monkeypatch):
    paths = fixture_sources(store, monkeypatch)
    class QuotaError(Exception):
        status_code, code = 429, 'insufficient_quota'
    class NoQuota(MockTagger):
        def tag(self, path, context=''):
            raise QuotaError('secret-key-material')
    monkeypatch.setattr(tasks, 'get_tagger', lambda _: NoQuota())
    job = tasks.create_import(paths, 'mock')
    tasks.run_job(job['id'])
    public = tasks.public_job(tasks.get_job(job['id']))
    assert public['status'] == 'failed' and public['processed'] == 1
    assert '额度不足' in public['message'] and 'secret-key-material' not in str(public)
    assert catalog.load_catalog() == []


def test_review_filter_and_locate_selected_on_later_page(store):
    records = [make_record(store, f'sweater-{i}') for i in range(35)]
    for i, record in enumerate(records):
        record.source.batch_id = 'batch-a'
        record.source.source_file = 'lookbook.pdf'
        if i == 34: record.proposed_tags.品类 = '外套'
    catalog.save_catalog(records)
    client = TestClient(api.app)
    response = client.get('/api/reviews', params={'selected_id': records[33].id}).json()
    assert response['page'] == 2 and response['items'][3]['id'] == records[33].id
    filtered = client.get('/api/reviews', params={'batch': 'batch-a', 'source_file': 'lookbook.pdf', 'category': '外套'}).json()
    assert filtered['count'] == 1 and filtered['items'][0]['id'] == records[-1].id
    assert client.get('/api/reviews', params={'batch': 'missing'}).json()['count'] == 0


def test_retag_job_preserves_human_tags_and_requires_new_review(store):
    record = make_record(store)
    catalog.save_catalog([record])
    approved = approve(record, record.tags.model_copy(update={'品类': '外套'}))
    job = tasks.create_retag('mock', True)
    tasks.run_job(job['id'])
    record = catalog.load_catalog()[0]
    assert record.tags.品类 == '外套' and record.proposed_tags.品类 == '毛衣'
    assert record.review_status == 'pending' and record.revision > approved.revision


def test_similar_records_excludes_self_and_unapproved_and_protects_source(store):
    first, other, pending = [make_record(store, name, color) for name, color in [('sweater','navy'),('coat','red'),('dress','white')]]
    catalog.save_catalog([first, other, pending])
    first, other = approve(first), approve(other)
    client = TestClient(api.app)
    hits = client.get(f'/api/records/{first.id}/similar').json()['items']
    assert [item['id'] for item in hits] == [other.id]
    first.source.source_path = str(store / 'private.txt')
    (store / 'private.txt').write_text('private')
    catalog.save_catalog([first])
    assert client.get(f'/api/records/{first.id}').json()['source_url'] is None
    assert client.get(f'/api/records/{first.id}/source').status_code == 404


def test_stale_retag_target_fails_before_paid_model_call(store, monkeypatch):
    record = make_record(store)
    catalog.save_catalog([record])
    job = tasks.create_retag('mock', True)
    approve(record)
    class MustNotCall(MockTagger):
        def tag(self, path, context=''):
            raise AssertionError('Stale target must be checked before invoking the model')
    monkeypatch.setattr(tasks, 'get_tagger', lambda _: MustNotCall())
    tasks.run_job(job['id'])
    public = tasks.public_job(tasks.get_job(job['id']))
    assert public['failed'] == 1 and '已被其他操作修改' in public['errors'][0]['message']
    assert catalog.load_catalog()[0].review_status == 'approved'


def test_immediate_retry_cannot_get_stuck_behind_finishing_worker(store, monkeypatch):
    paths = fixture_sources(store, monkeypatch, 1)
    calls, completed = [], Event()
    class OnceFail(MockTagger):
        def tag(self, path, context=''):
            calls.append(path)
            if len(calls) == 1: raise TimeoutError('transient')
            return super().tag(path, context)
    monkeypatch.setattr(tasks, 'get_tagger', lambda _: OnceFail())
    original_finish = tasks._finish
    def finish_then_retry(job, status, message, end=None):
        original_finish(job, status, message, end)
        if status == 'failed' and len(calls) == 1:
            tasks.retry_job(job['id'])
            tasks.submit_job(job['id'])
        elif status == 'completed': completed.set()
    monkeypatch.setattr(tasks, '_finish', finish_then_retry)
    job = tasks.create_import(paths, 'mock')
    try:
        tasks.submit_job(job['id'])
        assert completed.wait(5)
    finally:
        tasks.shutdown_worker()
    assert len(calls) == 2 and tasks.get_job(job['id'])['status'] == 'completed'
