"""Extract embedded images from PPTX / XLSX / PDF with source metadata."""

from __future__ import annotations

import hashlib
import io
import math
import zipfile
from pathlib import Path

import pymupdf
from openpyxl import load_workbook
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from src.catalog import EXTRACTED_DIR, ensure_data_dirs
from src.extract.pdf_notes import page_notes
from src.schema import DesignNotes, ImageRecord, SourceMeta

DOCUMENT_SUFFIXES = {".pptx", ".xlsx", ".xlsm", ".pdf"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}
SUPPORTED_SUFFIXES = DOCUMENT_SUFFIXES | IMAGE_SUFFIXES
# Skip tiny icons / logos commonly embedded in lookbooks.
MIN_IMAGE_SIDE = 120
MIN_IMAGE_PIXELS = 20_000
# PDF regions (points): merge tolerance, banner aspect cut-off, background page share.
PDF_TOLERANCE = 2.0
PDF_MAX_ASPECT = 3.5
PDF_MAX_PAGE_SHARE = 0.6
# Catalogue tables embed product shots as small as ~110x170 px.
PDF_MIN_SIDE = 96
PDF_MIN_PIXELS = 15_000
PDF_MAX_RENDER_PIXELS = 20_000_000
PDF_MAX_RENDER_SIDE = 8192


def _safe_stem(name: str) -> str:
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in name)[:80]


def _unique_path(directory: Path, stem: str, suffix: str, content: bytes) -> Path:
    digest = hashlib.md5(content).hexdigest()[:8]
    candidate = directory / f"{stem}_{digest}{suffix}"
    return candidate


def _save_image_bytes(
    content: bytes,
    out_dir: Path,
    stem: str,
    preferred_suffix: str = ".png",
) -> Path | None:
    try:
        image = Image.open(io.BytesIO(content))
        image.load()
    except Exception:
        return None

    if min(image.size) < MIN_IMAGE_SIDE or (image.size[0] * image.size[1]) < MIN_IMAGE_PIXELS:
        return None

    original_fmt = (image.format or "").upper()
    orientation = None
    # Bake EXIF orientation if present, then drop orientation tag.
    try:
        from PIL import ImageOps

        orientation = image.getexif().get(274, 1)
        image = ImageOps.exif_transpose(image) or image
    except Exception:
        pass

    suffix = {"JPEG": ".jpg", "PNG": ".png", "WEBP": ".webp",
              "GIF": ".gif", "BMP": ".bmp"}.get(original_fmt, ".png")
    if orientation != 1:
        suffix = ".png"  # Store corrected pixels losslessly, including rotated JPEGs.
    path = _unique_path(out_dir, stem, suffix, content)
    if path.exists():
        return path
    # Keep embedded photo bytes exactly when orientation needs no correction.
    # Re-saving JPEGs with Pillow defaults used to introduce a second lossy encode.
    if original_fmt in {"JPEG", "PNG", "WEBP", "GIF", "BMP"} and orientation == 1:
        path.write_bytes(content)
        return path

    if image.mode not in ("RGB", "RGBA"):
        image = image.convert("RGBA" if "A" in image.getbands() else "RGB")

    image.save(path, format="PNG")
    return path


def _stacked_groups(rects: list[pymupdf.Rect], tol: float = PDF_TOLERANCE) -> list[pymupdf.Rect]:
    """Union image tiles that share left/right edges and touch vertically.

    Lookbook exporters often slice one model photo into horizontal strips; side-by-side
    images are left alone because they are usually separate thumbnails.
    """
    parent = list(range(len(rects)))

    def find(i: int) -> int:
        while parent[i] != i:
            parent[i] = parent[parent[i]]
            i = parent[i]
        return i

    for i, a in enumerate(rects):
        for j in range(i + 1, len(rects)):
            b = rects[j]
            if (abs(a.x0 - b.x0) < tol and abs(a.x1 - b.x1) < tol
                    and min(abs(a.y1 - b.y0), abs(b.y1 - a.y0)) < tol):
                parent[find(i)] = find(j)
    merged: dict[int, pymupdf.Rect] = {}
    for i, rect in enumerate(rects):
        merged.setdefault(find(i), pymupdf.Rect(rect)).include_rect(rect)
    return list(merged.values())


def pdf_render_scale(region: pymupdf.Rect, infos: list[tuple[pymupdf.Rect, dict]]) -> float:
    """Preserve source sampling even for rotated, cropped or very small on-page photos."""
    densities = [1.0]
    for rect, info in infos:
        if not region.contains(rect):
            continue
        a, b, c, d, *_ = info["transform"]
        # Transform axes describe the full image, unlike the clipped bounding box.
        x_points, y_points = math.hypot(a, b), math.hypot(c, d)
        if x_points > 0 and y_points > 0:
            densities.append(max(info["width"] / x_points, info["height"] / y_points))
    return min(max(densities),
               math.sqrt(PDF_MAX_RENDER_PIXELS / (region.width * region.height)),
               PDF_MAX_RENDER_SIDE / max(region.width, region.height))


def extract_from_pdf(path: Path, out_dir: Path) -> list[ImageRecord]:
    """Render each on-page picture region, so sliced, flipped or masked images come out as shown."""
    records: list[ImageRecord] = []
    stem = _safe_stem(path.stem)
    seen: set[str] = set()
    with pymupdf.open(path) as doc:
        for page_idx, page in enumerate(doc):
            infos = [(pymupdf.Rect(info["bbox"]) & page.rect, info) for info in page.get_image_info()]
            infos = [(rect, info) for rect, info in infos if not rect.is_empty]
            page_area = page.rect.width * page.rect.height
            regions = _stacked_groups([rect for rect, _ in infos])
            notes = page_notes(page, regions)
            for image_idx, region in enumerate(regions, start=1):
                if min(region.width, region.height) < 12:
                    continue  # tiny layout icons must not become large photos at native resolution
                if max(region.width, region.height) / max(1, min(region.width, region.height)) > PDF_MAX_ASPECT:
                    continue  # banners and divider strips
                if region.width * region.height > PDF_MAX_PAGE_SHARE * page_area:
                    continue  # full-page backgrounds
                scale = pdf_render_scale(region, infos)
                pix = page.get_pixmap(clip=region, matrix=pymupdf.Matrix(scale, scale),
                                      colorspace=pymupdf.csRGB, alpha=False)
                if min(pix.width, pix.height) < PDF_MIN_SIDE or pix.width * pix.height < PDF_MIN_PIXELS:
                    continue
                digest = hashlib.md5(pix.samples).hexdigest()
                if digest in seen:
                    continue
                seen.add(digest)
                saved = out_dir / f"{stem}_page{page_idx + 1}_{image_idx:02d}_{digest[:8]}.png"
                saved.write_bytes(pix.tobytes("png"))
                records.append(ImageRecord(
                    image_path=str(saved.resolve()),
                    source=SourceMeta(source_file=path.name, source_type="pdf", location=f"page {page_idx + 1}",
                                      pdf_image_index=image_idx),
                    notes=notes.get(image_idx - 1, DesignNotes()),
                ))
    return records


def extract_from_pptx(path: Path, out_dir: Path) -> list[ImageRecord]:
    prs = Presentation(str(path))
    records: list[ImageRecord] = []
    stem = _safe_stem(path.stem)

    for slide_idx, slide in enumerate(prs.slides, start=1):
        image_idx = 0
        for shape in slide.shapes:
            if shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
                continue
            image_idx += 1
            image = shape.image
            content = image.blob
            ext = f".{(image.ext or 'png').lstrip('.')}"
            saved = _save_image_bytes(
                content,
                out_dir,
                f"{stem}_slide{slide_idx}_{image_idx:02d}",
                preferred_suffix=ext,
            )
            if saved is None:
                continue
            records.append(
                ImageRecord(
                    image_path=str(saved.resolve()),
                    source=SourceMeta(
                        source_file=path.name,
                        source_type="pptx",
                        location=f"slide {slide_idx}",
                    ),
                )
            )
    return records


def extract_from_xlsx(path: Path, out_dir: Path) -> list[ImageRecord]:
    """Extract images from xlsx; prefer sheet-level images so location keeps Sheet 名."""
    records: list[ImageRecord] = []
    stem = _safe_stem(path.stem)
    wb = load_workbook(str(path))
    image_idx = 0

    try:
        for sheet in wb.worksheets:
            images = list(getattr(sheet, "_images", []) or [])
            for img in images:
                try:
                    content = img._data()  # noqa: SLF001 - openpyxl private but stable enough for MVP
                except Exception:
                    continue
                image_idx += 1
                saved = _save_image_bytes(
                    content,
                    out_dir,
                    f"{stem}_{_safe_stem(sheet.title)}_{image_idx:02d}",
                    preferred_suffix=".png",
                )
                if saved is None:
                    continue
                records.append(
                    ImageRecord(
                        image_path=str(saved.resolve()),
                        source=SourceMeta(
                            source_file=path.name,
                            source_type="xlsx",
                            location=sheet.title,
                        ),
                    )
                )
    finally:
        wb.close()

    if records:
        return records

    # Fallback: zip media when sheet._images is empty (some exporters).
    with zipfile.ZipFile(path) as zf:
        sheet_names = []
        try:
            wb2 = load_workbook(str(path), read_only=True)
            sheet_names = list(wb2.sheetnames)
            wb2.close()
        except Exception:
            sheet_names = ["Sheet"]

        media_names = [
            name
            for name in zf.namelist()
            if name.startswith("xl/media/") and Path(name).suffix.lower() in IMAGE_SUFFIXES
        ]
        for idx, name in enumerate(sorted(media_names), start=1):
            content = zf.read(name)
            saved = _save_image_bytes(
                content,
                out_dir,
                f"{stem}_sheetimg_{idx:02d}",
                preferred_suffix=Path(name).suffix.lower(),
            )
            if saved is None:
                continue
            # Best-effort sheet label when drawing→sheet map is unavailable.
            sheet_guess = sheet_names[(idx - 1) % len(sheet_names)]
            records.append(
                ImageRecord(
                    image_path=str(saved.resolve()),
                    source=SourceMeta(
                        source_file=path.name,
                        source_type="xlsx",
                        location=f"{sheet_guess} / {Path(name).name}",
                    ),
                )
            )
    return records


def extract_from_image(path: Path, out_dir: Path) -> list[ImageRecord]:
    """One uploaded photo becomes one record. The caller sets 款号 from the filename."""
    saved = _save_uploaded_image(path.read_bytes(), out_dir, _safe_stem(path.stem) or "image")
    return [ImageRecord(
        image_path=str(saved.resolve()),
        source=SourceMeta(source_file=path.name, source_type="image", location="上传图片"),
    )]


def _save_uploaded_image(content: bytes, out_dir: Path, stem: str) -> Path:
    try:
        image = Image.open(io.BytesIO(content))
        original_fmt = (image.format or "").upper()
        image.load()
    except Exception as exc:
        raise ValueError("无法读取图片，请确认文件未损坏") from exc
    orientation = 1
    try:
        from PIL import ImageOps

        orientation = image.getexif().get(274, 1) or 1
        image = ImageOps.exif_transpose(image) or image
    except Exception:
        pass
    suffix = {"JPEG": ".jpg", "PNG": ".png", "WEBP": ".webp", "GIF": ".gif", "BMP": ".bmp"}.get(original_fmt, ".png")
    if orientation not in (1, None):
        suffix = ".png"
    path = _unique_path(out_dir, stem, suffix, content)
    if path.exists():
        return path
    if original_fmt in {"JPEG", "PNG", "WEBP", "GIF", "BMP"} and orientation in (1, None):
        path.write_bytes(content)
        return path
    if image.mode not in ("RGB", "RGBA"):
        image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
    image.save(path, format="PNG")
    return path


def extract_from_file(path: Path | str, out_dir: Path | None = None) -> list[ImageRecord]:
    ensure_data_dirs()
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(source)

    target = out_dir or (EXTRACTED_DIR / _safe_stem(source.stem))
    target.mkdir(parents=True, exist_ok=True)

    suffix = source.suffix.lower()
    if suffix == ".pptx":
        return extract_from_pptx(source, target)
    if suffix in {".xlsx", ".xlsm"}:
        return extract_from_xlsx(source, target)
    if suffix == ".pdf":
        return extract_from_pdf(source, target)
    if suffix in IMAGE_SUFFIXES:
        return extract_from_image(source, target)
    raise ValueError(f"Unsupported file type: {suffix}. Use .pptx, .xlsx, .pdf, or an image file")
