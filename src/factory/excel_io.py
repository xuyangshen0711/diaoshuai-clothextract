"""Shared Excel header detection for ERP exports.

Real uploads have a single header row. The wider sample puts a marker row
(重要 / 第二级别保留) above the real header; columns are found by name, not position.
"""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


def load_sheet_rows(path: Path) -> list[tuple]:
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("请上传 Excel 文件（.xlsx 或 .xlsm）")
    if not path.is_file():
        raise FileNotFoundError(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [tuple(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def header_text(value) -> str:
    if value is None:
        return ""
    text = unicodedata.normalize("NFKC", str(value)).replace("\n", "").replace("\r", "")
    return re.sub(r"\s+", "", text)


def plain_text(value) -> str:
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return unicodedata.normalize("NFKC", str(value)).strip()


def find_header(rows: list[tuple], required: str) -> tuple[int, list[str]]:
    for index, row in enumerate(rows):
        headers = [header_text(cell) for cell in row]
        if required in headers:
            return index, headers
    raise ValueError(f"找不到表头列「{required}」。请确认第一行（或表头行）包含该列名。")


def column_index(headers: list[str], name: str) -> int | None:
    return headers.index(name) if name in headers else None


def sheet_headers(rows: list[tuple], limit: int = 8) -> set[str]:
    names: set[str] = set()
    for row in rows[:limit]:
        for cell in row:
            text = header_text(cell)
            if text:
                names.add(text)
    return names


def table_kind_from_headers(names: set[str]) -> str:
    """'archive', 'sales', or '' — used only to choose an error or hint sentence."""
    sales = "款号" in names and "订单数量" in names
    archive = "大货款号" in names
    if sales and not archive:
        return "sales"
    if archive and not sales:
        return "archive"
    if sales and archive:
        return "sales" if "订单金额" in names and "原料名称" not in names else "archive"
    return ""


def workbook_table_kind(path: Path) -> str:
    try:
        rows = load_sheet_rows(path)
    except (OSError, ValueError):
        return ""
    return table_kind_from_headers(sheet_headers(rows))


def cell_at(row: tuple, index: int | None):
    if index is None or index >= len(row):
        return None
    return row[index]


def parse_number(value) -> float | None:
    if value is None or value == "" or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = plain_text(value).replace(",", "").replace("，", "")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def as_number(value: float):
    number = float(value)
    if number.is_integer():
        return int(round(number))
    return round(number, 2)
